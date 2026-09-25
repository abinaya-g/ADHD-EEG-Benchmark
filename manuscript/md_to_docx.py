"""Minimal Markdown -> DOCX converter for manuscript.md.
No pandoc available in this environment, so this hand-rolled converter
handles exactly the markdown subset used in manuscript.md: headers (#-###),
horizontal rules, pipe tables, bold/italic inline spans, numbered and
bulleted lists, ![caption](path) images, and plain paragraphs. Not a
general-purpose converter."""
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_inline_runs(paragraph, text):
    # Handle **bold** and *italic* (non-greedy, non-nested)
    pos = 0
    pattern = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            paragraph.add_run(token[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def parse_table_block(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = lines[i]
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        rows.append(cells)
        i += 1
    # remove the separator row (---|---|---)
    if len(rows) > 1 and all(re.match(r"^:?-+:?$", c) for c in rows[1]):
        rows.pop(1)
    return rows, i


def build_docx(md_path, out_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")
    base_dir = os.path.dirname(os.path.abspath(md_path))

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped == "":
            i += 1
            continue

        if stripped == "---":
            doc.add_page_break()
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            heading_level = min(max(level, 1), 4)
            h = doc.add_heading(level=heading_level)
            add_inline_runs(h, text)
            i += 1
            continue

        img_m = re.match(r"^!\[(.*)\]\((.*)\)$", stripped)
        if img_m:
            caption, rel_path = img_m.group(1), img_m.group(2)
            img_path = os.path.join(base_dir, rel_path)
            if os.path.isfile(img_path):
                doc.add_picture(img_path, width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p = doc.add_paragraph()
                add_inline_runs(p, f"[Missing figure file: {rel_path}]")
            cap_p = doc.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_run = cap_p.add_run(caption)
            cap_run.italic = True
            cap_run.font.size = Pt(10)
            i += 1
            continue

        if stripped.startswith("|"):
            rows, next_i = parse_table_block(lines, i)
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Light Grid Accent 1"
                for r_idx, row in enumerate(rows):
                    for c_idx in range(ncols):
                        cell_text = row[c_idx] if c_idx < len(row) else ""
                        cell = table.cell(r_idx, c_idx)
                        cell.paragraphs[0].text = ""
                        add_inline_runs(cell.paragraphs[0], cell_text)
                        if r_idx == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph("")
            i = next_i
            continue

        if re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, re.sub(r"^\d+\.\s", "", stripped))
            i += 1
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:])
            i += 1
            continue

        if stripped.startswith("**Table") or stripped.startswith("**Figure") or stripped.startswith("*Source:"):
            p = doc.add_paragraph()
            add_inline_runs(p, stripped)
            i += 1
            continue

        # plain paragraph (accumulate until blank line for nicer wrapping)
        para_lines = [stripped]
        j = i + 1
        while j < n and lines[j].strip() != "" and not lines[j].strip().startswith(("#", "|", "---", "![")) \
                and not re.match(r"^\d+\.\s", lines[j].strip()) and not lines[j].strip().startswith(("- ", "* ")):
            para_lines.append(lines[j].strip())
            j += 1
        p = doc.add_paragraph()
        add_inline_runs(p, " ".join(para_lines))
        i = j

    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    build_docx(sys.argv[1], sys.argv[2])
